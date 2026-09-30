from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
import subprocess
import sys
import tempfile

import pymupdf
import numpy as np
from PIL import Image
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.config import settings

SUPPORTED_SUFFIXES = {".md", ".markdown", ".pdf", ".docx", ".doc"}
_HORIZONTAL_WHITESPACE = re.compile(r"[^\S\n]+")
_FENCE = re.compile(r"^[ \t]{0,3}(`{3,}|~{3,})")
_OCR_ENGINE = None


@dataclass
class LoadedDocument:
    text: str
    metadata: dict = field(default_factory=dict)


def normalize_document_text(text: str) -> str:
    """Normalize document whitespace without changing fenced code blocks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized: list[str] = []
    in_fence = False
    fence_char = ""
    fence_length = 0
    blank_line_pending = False

    for line in text.split("\n"):
        fence_match = _FENCE.match(line)
        if fence_match:
            fence = fence_match.group(1)
            char = fence[0]
            length = len(fence)
            if not in_fence:
                in_fence = True
                fence_char = char
                fence_length = length
                normalized.append(line)
                blank_line_pending = False
                continue
            if char == fence_char and length >= fence_length:
                in_fence = False
                fence_char = ""
                fence_length = 0
                normalized.append(line)
                blank_line_pending = False
                continue

        if in_fence:
            normalized.append(line)
            continue

        cleaned = _HORIZONTAL_WHITESPACE.sub(" ", line).strip()
        if cleaned:
            normalized.append(cleaned)
            blank_line_pending = False
        elif normalized and not blank_line_pending:
            normalized.append("")
            blank_line_pending = True

    while normalized and not normalized[0]:
        normalized.pop(0)
    while normalized and not normalized[-1]:
        normalized.pop()

    return "\n".join(normalized)


def _table_to_markdown(table: Table) -> str:
    rows = []
    for row in table.rows:
        cells = [
            cell.text.replace("\n", " ").replace("|", "\\|").strip()
            for cell in row.cells
        ]
        rows.append(cells)

    # 1. cell.text 获取单元格原始文字
    # 2. replace("\n", " ")：单元格内的换行，替换成空格（避免markdown被意外换行打断）
    # 3. replace("|", "\\|")：转义竖线！单元格内容如果有 |，要写成 \|，否则破坏markdown表格语法
    # 4. strip()：去掉单元格首尾空白
    if not rows:
        return ""
    header = rows[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _paragraph_to_markdown(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""

    #根据标题转化为“#”
    #如 1-># ；1.1->## ; 1.1.1-->###
    style_name = (paragraph.style.name if paragraph.style else "") or ""
    if style_name.lower().startswith("heading"):
        suffix = style_name.split()[-1]
        level = int(suffix) if suffix.isdigit() else 1
        return f"{'#' * min(max(level, 1), 6)} {text}"
    return text


def load_markdown(path: Path) -> LoadedDocument:
    return LoadedDocument(
        text=path.read_text(encoding="utf-8"),
        metadata={"source": path.name, "type": "markdown"},
    )


def load_docx(path: Path) -> LoadedDocument:
    document = Document(str(path))
    parts: list[str] = []

    for child in document.element.body.iterchildren():      #遍历 docx 正文`<w:body>`标签的 所有直接子 XML 节点 ，逐个判断是段落还是表格，分别处理
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            markdown = _paragraph_to_markdown(paragraph)
            if markdown:
                parts.append(markdown)
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            markdown = _table_to_markdown(table)
            if markdown:
                parts.append(markdown)

    return LoadedDocument(
        text="\n\n".join(parts),
        metadata={"source": path.name, "type": "docx"},
    )


def _convert_doc_to_docx(path: Path, output_dir: Path) -> Path:
    output_path = output_dir / f"{path.stem}.docx"
    script = """
import sys
import pythoncom
import win32com.client

source, target = sys.argv[1], sys.argv[2]
pythoncom.CoInitialize()
word = None
document = None
try:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    document = word.Documents.Open(
        source,
        ReadOnly=True,
        AddToRecentFiles=False,
    )
    document.SaveAs2(target, FileFormat=16)
finally:
    if document is not None:
        document.Close(False)
    if word is not None:
        word.Quit()
    pythoncom.CoUninitialize()
"""
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(path.resolve()),
            str(output_path.resolve()),
        ],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if result.returncode != 0:
        detail = (result.stderr or result.stdout).strip()
        raise ValueError(f"无法转换 .doc 文件: {detail}")

    return output_path


def load_doc(path: Path) -> LoadedDocument:         # doc-->docx，再读取
    with tempfile.TemporaryDirectory(prefix="ihelp-doc-") as temp_dir:
        docx_path = _convert_doc_to_docx(path, Path(temp_dir))
        loaded = load_docx(docx_path)
    loaded.metadata["source"] = path.name
    loaded.metadata["type"] = "doc"
    return loaded


"""
    懒加载OCR引擎，全局单例。
    全局变量 _OCR_ENGINE 保存OCR实例，只初始化一次，避免反复加载ONNX模型。
"""
def _get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR

        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


"""
    对PDF的一页（fitz页对象）执行OCR识别，返回识别出来的文本字符串
    page: pymupdf（fitz）的PDF页对象
    return: OCR识别得到的文本
"""
def _ocr_page(page: pymupdf.Page) -> str:
    matrix = pymupdf.Matrix(settings.ocr_render_zoom, settings.ocr_render_zoom) #   构造缩放矩阵，ocr_render_zoom是配置项，比如2。放大图片，提升OCR识别精度
    pixmap = page.get_pixmap(matrix=matrix, alpha=False)        #    将PDF页面渲染成图片pixmap，alpha=False不要透明通道
    image = Image.open(BytesIO(pixmap.tobytes("png"))).convert("RGB")   # pixmap转png二进制，放进内存字节流BytesIO，再用PIL打开，转RGB图片

    # 调用OCR引擎，传入图片numpy数组
    # result是OCR输出列表，_丢弃额外信息
    result, _ = _get_ocr_engine()(np.array(image))
    if not result:
        return ""
    return "\n".join(   # 遍历OCR结果：item[1]是识别文字，过滤空文本，用换行拼接所有文字
        item[1].strip()
        for item in result
        if len(item) > 1 and item[1].strip()
    )


"""
    加载PDF文档，自动判断是否启用OCR；
    普通可复制文本直接提取，扫描件文字太少时自动走OCR。
    path：pdf文件路径Path对象
    use_ocr：是否开启ocr，None则读取全局配置settings.ocr_enabled
    返回 LoadedDocument 对象：包含完整文档文本 + metadata元信息
    
    但是没有对pdf做表格提取
"""
def load_pdf(path: Path, use_ocr: bool | None = None) -> LoadedDocument:
    document = pymupdf.open(str(path))
    parts: list[str] = []
    enable_ocr = settings.ocr_enabled if use_ocr is None else use_ocr

    for page_number, page in enumerate(document, 1):
        text = page.get_text("text").strip()
        if enable_ocr and len(text) < settings.ocr_min_text_chars:
            text = _ocr_page(page).strip()
        if text:
            parts.append(f"## 第 {page_number} 页\n\n{text}")

    return LoadedDocument(
        text="\n\n".join(parts),
        metadata={"source": path.name, "type": "pdf", "pages": len(document)},
    )


def load_document(path: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        return load_markdown(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix == ".doc":
        return load_doc(path)
    if suffix == ".pdf":
        return load_pdf(path)
    raise ValueError(f"不支持的文档类型: {suffix}")

import pymupdf
from docx import Document

from app.services import document_loader


def test_load_markdown(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("# 售后政策\n\n运费说明", encoding="utf-8")

    loaded = document_loader.load_document(path)

    assert "售后政策" in loaded.text
    assert loaded.metadata["type"] == "markdown"


def test_load_docx_preserves_heading_and_table(tmp_path):
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("售后政策", level=1)
    document.add_heading("运费说明", level=2)
    document.add_paragraph("普通商品满 99 元包邮。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "地区"
    table.cell(0, 1).text = "运费"
    table.cell(1, 0).text = "普通地区"
    table.cell(1, 1).text = "8 元"
    document.save(path)

    loaded = document_loader.load_document(path)

    assert "# 售后政策" in loaded.text
    assert "## 运费说明" in loaded.text
    assert "| 地区 | 运费 |" in loaded.text


def test_load_text_pdf(tmp_path):
    path = tmp_path / "policy.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Refund policy")
    document.save(path)
    document.close()

    loaded = document_loader.load_pdf(path, use_ocr=False)

    assert "Refund policy" in loaded.text
    assert loaded.metadata["pages"] == 1


def test_load_pdf_uses_ocr_when_text_is_too_short(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()

    monkeypatch.setattr(document_loader, "_ocr_page", lambda page: "OCR 识别文本")

    loaded = document_loader.load_pdf(path, use_ocr=True)

    assert "OCR 识别文本" in loaded.text


def test_normalize_document_text_collapses_whitespace_and_blank_lines():
    raw = "# 标题  \n\n  正文\t内容\u00a0\u3000 测试   \n   \n \n尾  部"

    normalized = document_loader.normalize_document_text(raw)

    assert normalized == "# 标题\n\n正文 内容 测试\n\n尾 部"


def test_normalize_document_text_preserves_fenced_code_block():
    raw = "说明  \n\n```python\n  value  =  1\n\n\n  print(value)\n```\n\n结束"

    normalized = document_loader.normalize_document_text(raw)

    assert (
        normalized
        == "说明\n\n```python\n  value  =  1\n\n\n  print(value)\n```\n\n结束"
    )
